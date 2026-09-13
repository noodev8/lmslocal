#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
Generates public/last-man-standing-template.xlsx - the free spreadsheet offered on
/last-man-standing-template.

WHY THE SCRIPT IS COMMITTED ALONGSIDE THE FILE IT MAKES. The team list goes stale every summer,
and a binary nobody can rebuild is a binary that quietly lies about who is in the Premier League.
Rerun this when the teams change:

    python scripts/make-lms-template.py

TEAMS COME FROM THE PRODUCT, NOT FROM MEMORY. The list below was taken from the live `team` table
(team_list_id 1, "English Premier League 2026-27") on 2026-09-13, so the sheet and the app agree
about who is playing. Refresh it with:

    cd lmslocal-server && node db/query.js "SELECT name FROM team WHERE team_list_id=1 ORDER BY name"

WHAT THE SHEET DELIBERATELY DOES NOT DO. It does not work out eliminations. That is not an
oversight and it is not a crippled-demo tactic - doing it properly in a spreadsheet needs a
results feed and it is exactly the job that makes people give up around round five. The sheet
automates the two things a spreadsheet is genuinely good at (counting teams used, catching a
player who has picked the same team twice) and leaves the status column honest and manual. The
page says so in as many words. Do not "improve" this into something that pretends otherwise.
"""

import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName

TEAMS = [
    'Arsenal', 'Aston Villa', 'Bournemouth', 'Brentford', 'Brighton', 'Chelsea', 'Coventry',
    'Crystal Palace', 'Everton', 'Fulham', 'Hull', 'Ipswich', 'Leeds', 'Liverpool', 'Man City',
    'Man Utd', 'Newcastle', 'Nottm Forest', 'Sunderland', 'Tottenham'
]

ROUNDS = 15          # A competition rarely runs past this; the sheet stays printable.
PLAYER_ROWS = 40

INK = '1C2620'       # Matches the site's ink token, so a printed sheet looks related to the site.
OVERPRINT = 'C0392B'
STOCK = 'EDEBE4'

thin = Side(style='thin', color='BFBCB2')
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)


def style_header(cell):
    cell.font = Font(bold=True, color='FFFFFF', size=10)
    cell.fill = PatternFill('solid', fgColor=INK)
    cell.alignment = Alignment(horizontal='center', vertical='center')
    cell.border = BORDER


def build():
    wb = Workbook()

    # ---------------------------------------------------------------- Picks
    ws = wb.active
    ws.title = 'Picks'

    ws['A1'] = 'LAST MAN STANDING'
    ws['A1'].font = Font(bold=True, size=18, color=INK)
    ws['A2'] = 'One team each round. A draw or a defeat and you are out. No team twice.'
    ws['A2'].font = Font(size=10, italic=True, color='5A6560')

    header_row = 4
    headers = ['Player'] + ['R%d' % n for n in range(1, ROUNDS + 1)] + ['Teams used', 'Check', 'Status']
    for i, name in enumerate(headers, start=1):
        c = ws.cell(row=header_row, column=i, value=name)
        style_header(c)

    first = header_row + 1
    last = header_row + PLAYER_ROWS
    pick_first_col = 2
    pick_last_col = 1 + ROUNDS
    used_col = pick_last_col + 1
    check_col = used_col + 1
    status_col = check_col + 1

    pl = get_column_letter(pick_first_col)
    pr = get_column_letter(pick_last_col)

    for r in range(first, last + 1):
        rng = '%s%d:%s%d' % (pl, r, pr, r)
        ws.cell(row=r, column=1).border = BORDER
        for col in range(pick_first_col, pick_last_col + 1):
            ws.cell(row=r, column=col).border = BORDER

        # How many teams they have spent. The count that tells a player when the pool resets.
        used = ws.cell(row=r, column=used_col, value='=COUNTA(%s)' % rng)
        used.alignment = Alignment(horizontal='center')
        used.border = BORDER

        # The single commonest scoring error: the same team picked twice. Blanks are excluded by
        # the (<>"") factor, so an empty row stays quiet instead of shouting CHECK at you.
        check = ws.cell(
            row=r, column=check_col,
            value='=IF(COUNTA({r})=0,"",IF(SUMPRODUCT((COUNTIF({r},{r})>1)*({r}<>""))>0,"CHECK","ok"))'.format(r=rng)
        )
        check.alignment = Alignment(horizontal='center')
        check.border = BORDER

        st = ws.cell(row=r, column=status_col)
        st.alignment = Alignment(horizontal='center')
        st.border = BORDER

    # Pick cells choose from the Teams sheet, so a typo cannot quietly invent a club.
    #
    # Via a DEFINED NAME rather than a direct Teams!$A$2:$A$41 reference, and with no leading "=".
    # Both matter: older Excel refuses a cross-sheet range in a validation list unless it is named,
    # and a formula1 beginning with "=" is written into the XML verbatim and rejected on open. A
    # broken dropdown is invisible until somebody downloads the file, which is too late.
    wb.defined_names.add(DefinedName('TeamList', attr_text="Teams!$A$2:$A$41"))
    dv_team = DataValidation(type='list', formula1='TeamList', allow_blank=True)
    dv_team.error = 'Pick a team from the list on the Teams sheet.'
    dv_team.errorTitle = 'Not one of the teams'
    ws.add_data_validation(dv_team)
    dv_team.add('%s%d:%s%d' % (pl, first, pr, last))

    dv_status = DataValidation(type='list', formula1='"IN,OUT"', allow_blank=True)
    ws.add_data_validation(dv_status)
    sc = get_column_letter(status_col)
    dv_status.add('%s%d:%s%d' % (sc, first, sc, last))

    # Visual cues: an eliminated row fades, a duplicate pick goes red.
    body = 'A%d:%s%d' % (first, sc, last)
    ws.conditional_formatting.add(
        body,
        FormulaRule(formula=['$%s%d="OUT"' % (sc, first)],
                    font=Font(color='9AA0A6', strike=True),
                    fill=PatternFill('solid', fgColor=STOCK))
    )
    cc = get_column_letter(check_col)
    ws.conditional_formatting.add(
        '%s%d:%s%d' % (cc, first, cc, last),
        FormulaRule(formula=['$%s%d="CHECK"' % (cc, first)],
                    font=Font(bold=True, color='FFFFFF'),
                    fill=PatternFill('solid', fgColor=OVERPRINT))
    )

    ws.column_dimensions['A'].width = 22
    for col in range(pick_first_col, pick_last_col + 1):
        ws.column_dimensions[get_column_letter(col)].width = 14
    ws.column_dimensions[get_column_letter(used_col)].width = 11
    ws.column_dimensions[get_column_letter(check_col)].width = 8
    ws.column_dimensions[sc].width = 9
    ws.freeze_panes = 'B%d' % first

    # ---------------------------------------------------------------- Teams
    t = wb.create_sheet('Teams')
    t['A1'] = 'Teams'
    style_header(t['A1'])
    for i, name in enumerate(TEAMS, start=2):
        t.cell(row=i, column=1, value=name).border = BORDER
    t['C2'] = 'Edit this list to match your league.'
    t['C3'] = 'The dropdowns on the Picks sheet read rows 2 to 41,'
    t['C4'] = 'so there is room to add more without changing anything.'
    for row in (2, 3, 4):
        t.cell(row=row, column=3).font = Font(size=10, italic=True, color='5A6560')
    t.column_dimensions['A'].width = 22
    t.column_dimensions['C'].width = 52

    # ---------------------------------------------------------------- Results
    res = wb.create_sheet('Results')
    res['A1'] = 'What happened'
    res['A1'].font = Font(bold=True, size=14, color=INK)
    res['A2'] = 'Your own record of each round. Nothing else reads this sheet.'
    res['A2'].font = Font(size=10, italic=True, color='5A6560')
    for i, name in enumerate(['Round', 'Date', 'Who went out', 'Still in', 'Notes'], start=1):
        style_header(res.cell(row=4, column=i, value=name))
    for r in range(5, 5 + ROUNDS):
        res.cell(row=r, column=1, value=r - 4).alignment = Alignment(horizontal='center')
        for c in range(1, 6):
            res.cell(row=r, column=c).border = BORDER
    for col, width in zip('ABCDE', (8, 14, 30, 12, 46)):
        res.column_dimensions[col].width = width

    # ---------------------------------------------------------------- Rules
    rules = wb.create_sheet('Rules')
    rules['A1'] = 'THE RULES'
    rules['A1'].font = Font(bold=True, size=18, color=INK)
    lines = [
        '',
        'Each round, every player picks one team they think will win.',
        'If that team wins, the player goes through to the next round.',
        'If the team draws or loses, the player is out.',
        '',
        'A player cannot pick the same team twice. Once every team has been used,',
        'they all become available again and the competition carries on.',
        '',
        'Results are judged on regulation time only - ninety minutes plus stoppage.',
        'Extra time and penalties do not count.',
        '',
        'A player who forgets to pick is treated as having lost that round. Decide before',
        'you start whether that means elimination or the loss of a life, and tell everybody.',
        '',
        'LIVES (optional)',
        'Giving each player a life or two means one bad week does not end their run.',
        'It makes a competition last longer, which is usually what you want with a small group.',
        '',
        'WHEN PICKS LOCK',
        'Set a deadline before the first kick-off of the round and hold to it.',
        'Late picks are the thing that causes arguments, not bad luck.',
        '',
        'IF EVERYONE GOES OUT IN THE SAME ROUND',
        'Common enough to decide in advance. Either the prize is shared between everyone',
        'eliminated in that final round, or the round is replayed with the same players.',
        '',
        '---',
        'Free sheet from LMSLocal - lmslocal.co.uk',
        'If the admin stops being fun, the site does all of the above for you.',
    ]
    for i, line in enumerate(lines, start=2):
        c = rules.cell(row=i, column=1, value=line)
        if line.isupper() and line not in ('---',):
            c.font = Font(bold=True, size=11, color=OVERPRINT)
        else:
            c.font = Font(size=11, color=INK)
    rules.column_dimensions['A'].width = 92

    out = os.path.join(os.path.dirname(__file__), '..', 'public', 'last-man-standing-template.xlsx')
    out = os.path.abspath(out)
    wb.save(out)
    print('wrote %s (%d bytes)' % (out, os.path.getsize(out)))


if __name__ == '__main__':
    build()
